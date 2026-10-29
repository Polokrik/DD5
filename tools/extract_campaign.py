"""Extraction en lecture seule avec adresses de cellules/paragraphes/pages.

Usage: python3 tools/extract_campaign.py > /tmp/dd5-sources.json
Prérequis hors runtime du jeu: openpyxl, python-docx et pypdf.
Les PDF des modules officiels sont des références, pas un journal de partie.
"""
import json
from pathlib import Path
from openpyxl import load_workbook
from docx import Document
from pypdf import PdfReader

root = Path(__file__).resolve().parents[1]
sources = {}
for path in sorted(root.iterdir()):
    records = []
    if path.suffix == '.xlsx':
        book = load_workbook(path, read_only=True, data_only=True)
        for sheet in book:
            for row in sheet:
                cells = [f'{c.column_letter}: {c.value}' for c in row if c.value is not None]
                if cells:
                    records.append({'ref': f'{sheet.title}!{next(c.row for c in row if c.value is not None)}', 'text': ' | '.join(cells)})
        book.close()
    elif path.suffix == '.docx':
        doc = Document(path)
        records.extend({'ref': f'p{i}', 'text': p.text} for i, p in enumerate(doc.paragraphs, 1) if p.text.strip())
        for ti, table in enumerate(doc.tables, 1):
            records.extend({'ref': f'table{ti}/row{ri}', 'text': ' | '.join(c.text for c in row.cells)} for ri, row in enumerate(table.rows, 1))
    elif path.suffix == '.pdf' and not path.name.startswith(('D&D', 'D&D5')):
        records.extend({'ref': f'page{i}', 'text': page.extract_text() or ''} for i, page in enumerate(PdfReader(path).pages, 1))
    else:
        continue
    sources[path.name] = records
print(json.dumps(sources, ensure_ascii=False, indent=2))
